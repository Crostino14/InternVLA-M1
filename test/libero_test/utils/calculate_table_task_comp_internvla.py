"""
calculate_table_task_comp_internvla.py  — Task Composition L1/L2
"""

import re, math, argparse, os, glob
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


# ─── TASK ORDER & REFERENCE ────────────────────────────────────────────────────

def normalize(s): return " ".join(s.strip().lower().split())

def get_task_comp_l1_order():
    return [
        "Put the plate on the top of the cabinet",
        "Put the plate on the stove",
        "Put the cream cheese on the top of the cabinet",
        "Put the cream cheese on the plate",
        "Open the top layer of the drawer and put the cream cheese inside",
    ]

def get_task_comp_l2_order():
    return [
        "Open the middle drawer of the cabinet",
        "Put the bowl on the stove",
        "Put the cream cheese on the bowl and put the bowl on the plate",
        "Push the plate to the front of the stove and put the bowl on the plate",
        "Put the cream cheese on the bowl and put the bowl on the top of the cabinet",
    ]

def get_reference_mapping(level="l1"):
    if level == "l2":
        return {
            normalize("Open the middle drawer of the cabinet"):
                "Open the middle drawer of the cabinet  (L2)",
            normalize("Put the bowl on the stove"):
                "Put the bowl on the stove  /  Turn on the stove",
            normalize("Put the cream cheese on the bowl and put the bowl on the plate"):
                "Put the cream cheese in the bowl  /  Put the bowl on the plate",
            normalize("Push the plate to the front of the stove and put the bowl on the plate"):
                "Push the plate to the front of the stove  /  Put the bowl on the plate",
            normalize("Put the cream cheese on the bowl and put the bowl on the top of the cabinet"):
                "Put the cream cheese in the bowl  /  Put the bowl on top of the cabinet",
        }
    return {
        normalize("Put the plate on the top of the cabinet"):
            "Put the bowl on the top of the cabinet",
        normalize("Put the plate on the stove"):
            "Put the bowl on the stove",
        normalize("Put the cream cheese on the top of the cabinet"):
            "Put the wine bottle on the top of the cabinet",
        normalize("Put the cream cheese on the plate"):
            "Put the cream cheese in the bowl",
        normalize("Open the top layer of the drawer and put the cream cheese inside"):
            "Open the top layer of the drawer and put the bowl inside",
    }


# ─── PARSER ────────────────────────────────────────────────────────────────────

def parse_task_comp_log(filepath: str, num_trials: int = 50) -> list:
    """
    Parsa un file task comp.

    Righe chiave:
      TASK X/N (Task Composition L2)
      Command:   Put the cream cheese on the bowl ...
      Task SR: 0.0000 (0.0%)   ← DECIMALE con parentesi → ×100 per %
    """
    tasks, current = [], {}

    with open(filepath, "r", errors="replace") as f:
        lines = [l.rstrip("\r\n") for l in f]

    for line in lines:
        m = re.match(r"^TASK\s+(\d+)/\d+\s+\(Task Composition", line)
        if m:
            if current and "task_sr" in current:
                tasks.append(current)
            current = {"task_id": int(m.group(1))}
            continue

        if line.startswith("Command:") and current is not None:
            current["command"] = line.split("Command:", 1)[1].strip()
            continue

        # "Task SR: 0.0000 (0.0%)"  ← con parentesi
        m2 = re.match(r"^Task SR:\s*([0-9.]+)\s*\(", line)
        if m2 and current:
            sr_decimal          = float(m2.group(1))          # es. 0.0000
            current["task_sr"]  = sr_decimal * 100            # → 0.0  (percentuale)
            current["episodes"] = num_trials
            current["successes"] = int(round(sr_decimal * num_trials))

    if current and "task_sr" in current:
        tasks.append(current)

    tasks.sort(key=lambda t: t.get("task_id", 0))
    return tasks


def merge_task_comp_files(filepaths: list, num_trials: int = 50) -> tuple:
    task_rates, task_eps, all_keys = defaultdict(list), {}, []

    for fp in filepaths:
        for t in parse_task_comp_log(fp, num_trials):
            key = normalize(t.get("command", ""))
            if key not in task_rates:
                all_keys.append(key)
            task_rates[key].append(t["task_sr"])   # già in %
            task_eps[key] = num_trials

    merged = {k: sum(v)/len(v) for k, v in task_rates.items()}
    return merged, task_eps, all_keys


# ─── EXCEL ─────────────────────────────────────────────────────────────────────

def write_excel_task_comp(output_xlsx, txt_files_by_seed,
                           level="l1", num_trials=50, model_name="InternVLA-M1"):

    print("\n[INFO] Parsing log files (task comp)...")
    all_merged = []
    for seed_idx in range(3):
        fps = txt_files_by_seed[seed_idx]
        print(f"  Seed {seed_idx}: {len(fps)} file(s)")
        merged_rates, task_eps, all_keys = merge_task_comp_files(fps, num_trials)
        all_merged.append((merged_rates, task_eps, all_keys))
        print(f"    → {len(merged_rates)} task trovati")

    ref_mapping = get_reference_mapping(level)
    fixed_order = get_task_comp_l2_order() if level == "l2" else get_task_comp_l1_order()

    wb = Workbook()
    ws = wb.active
    ws.title = f"Task Comp {level.upper()}"

    headers = [
        "# Task", "Task Command", "Reference Task",
        "SR% — Seed 0", "Comp. Seed 0",
        "SR% — Seed 1", "Comp. Seed 1",
        "SR% — Seed 2", "Comp. Seed 2",
        "Mean SR% ± Std%", "Mean Comp.",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font      = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    all_seed_rates = [[], [], []]
    all_seed_comps = [[], [], []]

    for task_idx, task_cmd in enumerate(fixed_order, start=1):
        tk       = normalize(task_cmd)
        ref_task = ref_mapping.get(tk, "")
        seed_rates, seed_comps = [], []

        for seed_idx, (merged_rates, task_eps, _) in enumerate(all_merged):
            rate = merged_rates.get(tk, float("nan"))   # già in %
            ep   = task_eps.get(tk, num_trials)
            succ = int(round(rate / 100 * ep)) if not math.isnan(rate) else 0
            seed_rates.append(rate)
            seed_comps.append(f"{succ}/{ep}")
            if not math.isnan(rate):
                all_seed_rates[seed_idx].append(rate)
                all_seed_comps[seed_idx].append((succ, ep))

        valid = [r for r in seed_rates if not math.isnan(r)]
        if valid:
            mean_r = sum(valid) / len(valid)
            std_r  = math.sqrt(sum((r-mean_r)**2 for r in valid) / max(len(valid)-1, 1))
            mean_display = f"{mean_r:.1f}% ± {std_r:.1f}%"
            avg_succ     = int(round(mean_r / 100 * num_trials))
        else:
            mean_display, avg_succ = "N/A", 0

        ws.append([
            task_idx, task_cmd, ref_task,
            f"{seed_rates[0]:.1f}%" if not math.isnan(seed_rates[0]) else "N/A",
            seed_comps[0],
            f"{seed_rates[1]:.1f}%" if not math.isnan(seed_rates[1]) else "N/A",
            seed_comps[1],
            f"{seed_rates[2]:.1f}%" if not math.isnan(seed_rates[2]) else "N/A",
            seed_comps[2],
            mean_display,
            f"{avg_succ}/{num_trials}",
        ])

    # ── Riga finale ──
    final_row = ["", "Mean SR% ± Std%", ""]
    for seed_idx in range(3):
        rates = all_seed_rates[seed_idx]
        if rates:
            m = sum(rates)/len(rates)
            s = math.sqrt(sum((r-m)**2 for r in rates) / max(len(rates)-1, 1))
            ts = sum(c[0] for c in all_seed_comps[seed_idx])
            te = sum(c[1] for c in all_seed_comps[seed_idx])
            final_row.extend([f"{m:.2f}% ± {s:.2f}%", f"{ts}/{te}"])
        else:
            final_row.extend(["N/A", "0/0"])

    seed_means = [sum(all_seed_rates[i])/len(all_seed_rates[i])
                  for i in range(3) if all_seed_rates[i]]
    if seed_means:
        gm = sum(seed_means)/len(seed_means)
        gs = math.sqrt(sum((m-gm)**2 for m in seed_means) / max(len(seed_means)-1, 1))
        gt_s = sum(sum(c[0] for c in all_seed_comps[i]) for i in range(3))
        gt_e = sum(sum(c[1] for c in all_seed_comps[i]) for i in range(3))
        final_row.extend([f"{gm:.2f}% ± {gs:.2f}%", f"{gt_s}/{gt_e}"])
    else:
        final_row.extend(["N/A", "0/0"])

    ws.append(final_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    _autowidth(ws)
    ws.row_dimensions[1].height = 30
    wb.save(output_xlsx)
    print(f"\n[OK] Salvato: {output_xlsx}")


# ─── AUTO-DETECT ───────────────────────────────────────────────────────────────

def find_task_comp_files_by_seed(base_dir, level, pattern_prefix=None):
    if pattern_prefix is None:
        pattern_prefix = f"EVAL-task_comp_{level}-internvla_m1"
    result = {0: [], 1: [], 2: []}
    for seed_idx in range(3):
        matches = sorted(set(
            m for pat in [
                os.path.join(base_dir, f"{pattern_prefix}*seed{seed_idx}*.txt"),
                os.path.join(base_dir, f"{pattern_prefix}*seed_{seed_idx}*.txt"),
            ] for m in glob.glob(pat)
        ))
        result[seed_idx] = matches
        print(f"  {'✓' if matches else '✗'} Seed {seed_idx}: {len(matches)} file(s)")
    return result if all(result.values()) else None


def _autowidth(ws):
    for col in ws.columns:
        letter  = col[0].column_letter
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[letter].width = min(max_len + 2, 65)


# ─── MAIN ──────────────────────────────────────────────────────────────────────

def main():
    p = argparse.ArgumentParser()
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--txt_dir")
    g.add_argument("--manual", action="store_true")
    p.add_argument("--level",      default="l1", choices=["l1","l2"])
    p.add_argument("--pattern",    default=None)
    p.add_argument("--model_name", default="InternVLA-M1")
    p.add_argument("--num_trials", type=int, default=50)
    p.add_argument("--seed0", nargs="+")
    p.add_argument("--seed1", nargs="+")
    p.add_argument("--seed2", nargs="+")
    p.add_argument("output_xlsx")
    args = p.parse_args()

    if args.manual:
        if not (args.seed0 and args.seed1 and args.seed2):
            exit("[ERROR] --manual richiede --seed0 --seed1 --seed2")
        files = {0: args.seed0, 1: args.seed1, 2: args.seed2}
    else:
        files = find_task_comp_files_by_seed(args.txt_dir, args.level, args.pattern)
        if files is None:
            exit("[ERROR] File mancanti per alcuni seed!")

    write_excel_task_comp(args.output_xlsx, files,
                          level=args.level,
                          num_trials=args.num_trials,
                          model_name=args.model_name)

if __name__ == "__main__":
    main()
