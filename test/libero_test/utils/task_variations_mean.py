#!/usr/bin/env python3
"""
Generate per-task, per-level syntactic evaluation tables in Excel.

Output: one Excel file with 3 sheets (L1, L2, L3).
Each sheet contains stacked tables — one per task (0-9).
Each table has:
  Rows  : configured model/variant combinations
  Cols  : Model | Original (syn_base) | V1 | V2 | V3 | Mean SR%
  Footer: variation command strings aligned under each version column

Values are Mean SR% ± Std% across seeds (ddof=1).

Usage:
    python generate_task_tables.py
    python generate_task_tables.py --output_xlsx task_tables.xlsx
"""

import os, re, glob, argparse
import numpy as np
from collections import defaultdict
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ──────────────────────────────────────────────────────────────────────────────
# CONFIGURATION — edit this section to add/remove models
# ──────────────────────────────────────────────────────────────────────────────
BASE_ROOT = (
    "/mnt/beegfs/a.cardamone7/outputs/results/libero_goal/syntactic_command_variation/internvla-m1_l3_finetuned"
)

# Each entry: (display_name, path_containing_l1/ l2/ l3/ subdirs)
MODELS = [
    ("InternVLA-M1 10 eps", f"{BASE_ROOT}/10eps_50k_steps"),
    ("InternVLA-M1 25 eps", f"{BASE_ROOT}/25eps_50k_steps"),
    ("InternVLA-M1 50 eps", f"{BASE_ROOT}/50eps_50k_steps"),
    # ── add other models here ──────────────────────────────────────────────
    # ("OpenVLA-OFT",  f"{BASE_ROOT}/openvla-oft/<variant_dir>"),
    # ("TinyVLA",      f"{BASE_ROOT}/tinyvla/<variant_dir>"),
]

LEVELS       = ["l1", "l2", "l3"]
LEVEL_LABELS = {"l1": "L1", "l2": "L2", "l3": "L3"}
N_TASKS      = 10

# Canonical version order (extras appended sorted)
CANON_VERSIONS = ["syn_base", "v1", "v2", "v3"]
VERSION_HEADER = {
    "syn_base": "Original\nMean SR% ± Std%",
    "v1":       "V1\nMean SR% ± Std%",
    "v2":       "V2\nMean SR% ± Std%",
    "v3":       "V3\nMean SR% ± Std%",
}

# ──────────────────────────────────────────────────────────────────────────────
# STYLES — defined once, reused everywhere
# ──────────────────────────────────────────────────────────────────────────────
TITLE_FONT  = Font(name="Calibri", size=11, bold=True,  color="FFFFFF")
HDR_FONT    = Font(name="Calibri", size=10, bold=True,  color="FFFFFF")
MEAN_HDR_F  = Font(name="Calibri", size=10, bold=True,  color="FFFFFF")
DATA_FONT   = Font(name="Calibri", size=10)
BOLD_DATA_F = Font(name="Calibri", size=10, bold=True)
CMD_FONT    = Font(name="Calibri", size=9,  italic=True, color="555555")
NA_FONT     = Font(name="Calibri", size=10, italic=True, color="AAAAAA")
FOOTER_FONT = Font(name="Calibri", size=8,  italic=True, color="999999")

TITLE_FILL  = PatternFill("solid", fgColor="1F3864")
HDR_FILL    = PatternFill("solid", fgColor="2E75B6")
MEAN_HDR_FL = PatternFill("solid", fgColor="1A4F72")
ROW_A_FILL  = PatternFill("solid", fgColor="EBF3FB")
ROW_B_FILL  = PatternFill("solid", fgColor="FFFFFF")
CMD_FILL    = PatternFill("solid", fgColor="F2F2F2")

CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP   = Alignment(horizontal="left",   vertical="center", wrap_text=True, indent=1)
CENTER      = Alignment(horizontal="center", vertical="center")

THIN = Side(style="thin",   color="B0C4DE")
MED  = Side(style="medium", color="1F3864")

def tb():
    return Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def mb():
    return Border(left=MED, right=MED, top=MED, bottom=MED)

# ──────────────────────────────────────────────────────────────────────────────
# PARSING
# ──────────────────────────────────────────────────────────────────────────────
_VER_SR_RE  = re.compile(
    r"VERSION\s+(\S+)\s+RESULTS:.*?Success Rate:\s*([\d.]+)%", re.DOTALL)
_TASK_SR_RE = re.compile(r"Task SR:\s*([\d.]+)")
_ORIG_RE    = re.compile(r"Original Command:\s*(.+)")
_VAR_RE     = re.compile(r"Variation Command:\s*(.+)")
_SEP_RE     = re.compile(r"={8,}")
_ST_RE      = re.compile(r"seed(\d+)_task(\d+)")


def parse_file(filepath):
    """
    Returns:
        task_name      : str   — natural language task description
        version_srs    : dict  {version_key: float}  SR in % (0–100)
        task_sr        : float | None                overall Task SR in %
        variation_cmds : dict  {version_key: str}
    """
    with open(filepath, "r", errors="replace") as f:
        content = f.read()

    m = _ORIG_RE.search(content)
    task_name = m.group(1).strip() if m else "Unknown"

    version_srs = {
        m.group(1): float(m.group(2))
        for m in _VER_SR_RE.finditer(content)
    }

    m = _TASK_SR_RE.search(content)
    task_sr = float(m.group(1)) * 100.0 if m else None

    variation_cmds = {}
    for block in _SEP_RE.split(content):
        mv = re.search(r"Testing VERSION:\s*(\S+)", block)
        mc = _VAR_RE.search(block)
        if mv and mc:
            variation_cmds[mv.group(1)] = mc.group(1).strip()

    return task_name, version_srs, task_sr, variation_cmds


def collect_task_data(model_dir, level, task_id):
    """
    Reads all seed files for (model_dir, level, task_id).
    Deduplicates by (seed, task) key — keeps latest timestamp.
    Returns:
        task_name      : str | None
        version_stats  : dict {ver: (mean_sr, std_sr)}  in %
        task_sr_stats  : (mean_sr, std_sr) | (None, None)
        variation_cmds : dict {ver: str}
    """
    level_dir = os.path.join(model_dir, level)
    if not os.path.isdir(level_dir):
        return None, {}, (None, None), {}

    files = sorted(glob.glob(
        os.path.join(level_dir, f"*_task{task_id}-{level}.txt")))

    # Deduplicate by (seed, task): alphabetically last = newest wins
    seen = {}
    for fp in files:
        m = _ST_RE.search(os.path.basename(fp))
        key = (int(m.group(1)), int(m.group(2))) if m else fp
        seen[key] = fp

    if not seen:
        return None, {}, (None, None), {}

    task_name      = None
    ver_lists      = defaultdict(list)
    task_sr_list   = []
    variation_cmds = {}

    for fp in seen.values():
        t, v_srs, t_sr, v_cmds = parse_file(fp)
        if task_name is None:
            task_name = t
        for ver, sr in v_srs.items():
            ver_lists[ver].append(sr)
        if t_sr is not None:
            task_sr_list.append(t_sr)
        if not variation_cmds and v_cmds:
            variation_cmds = v_cmds

    version_stats = {}
    for ver, vals in ver_lists.items():
        a = np.array(vals)
        version_stats[ver] = (
            np.mean(a),
            np.std(a, ddof=1) if len(a) > 1 else 0.0
        )

    if task_sr_list:
        a = np.array(task_sr_list)
        task_sr_stats = (
            np.mean(a),
            np.std(a, ddof=1) if len(a) > 1 else 0.0
        )
    else:
        task_sr_stats = (None, None)

    return task_name, version_stats, task_sr_stats, variation_cmds


# ──────────────────────────────────────────────────────────────────────────────
# EXCEL BUILDER
# ──────────────────────────────────────────────────────────────────────────────
def _fmt(mean, std):
    return f"{mean:.1f}% ± {std:.1f}%"


def write_task_table(ws, start_row, level_label, task_id, task_name,
                     models_data, versions):
    """
    Write one task table into ws starting at start_row.

    models_data : list of (model_name, version_stats, task_sr_stats, var_cmds)
    versions    : ordered list of version keys present in this table
    Returns     : next available row (after 1-row gap)
    """
    COL_MODEL = 2
    COL_V0    = 3
    n_vers    = len(versions)
    COL_MEAN  = COL_V0 + n_vers
    LAST_COL  = COL_MEAN

    # ── Title ────────────────────────────────────────────────────────────────
    ws.merge_cells(start_row=start_row, start_column=COL_MODEL,
                   end_row=start_row,   end_column=LAST_COL)
    c = ws.cell(row=start_row, column=COL_MODEL,
                value=(f'{level_label} Syntactic Variation Evaluation — '
                       f'Task {task_id}: "{task_name}"'))
    c.font = TITLE_FONT; c.fill = TITLE_FILL; c.alignment = LEFT_WRAP
    ws.row_dimensions[start_row].height = 22

    # ── Header ───────────────────────────────────────────────────────────────
    hr = start_row + 1
    ws.row_dimensions[hr].height = 30

    c = ws.cell(row=hr, column=COL_MODEL, value="Model")
    c.font = HDR_FONT; c.fill = HDR_FILL
    c.alignment = CENTER_WRAP; c.border = tb()

    for i, ver in enumerate(versions):
        label = VERSION_HEADER.get(ver, f"{ver}\nMean SR% ± Std%")
        c = ws.cell(row=hr, column=COL_V0 + i, value=label)
        c.font = HDR_FONT; c.fill = HDR_FILL
        c.alignment = CENTER_WRAP; c.border = tb()

    c = ws.cell(row=hr, column=COL_MEAN, value="Mean SR%\n± Std%")
    c.font = MEAN_HDR_F; c.fill = MEAN_HDR_FL
    c.alignment = CENTER_WRAP; c.border = tb()

    # ── Data rows ────────────────────────────────────────────────────────────
    fills = [ROW_A_FILL, ROW_B_FILL, ROW_A_FILL, ROW_B_FILL]
    collected_cmds = {}

    for i, (model_name, ver_stats, task_sr_stats, var_cmds) in enumerate(models_data):
        dr   = hr + 1 + i
        fill = fills[i % len(fills)]
        ws.row_dimensions[dr].height = 17

        c = ws.cell(row=dr, column=COL_MODEL, value=model_name)
        c.font = BOLD_DATA_F; c.fill = fill
        c.alignment = LEFT_WRAP; c.border = tb()

        for j, ver in enumerate(versions):
            col = COL_V0 + j
            if ver in ver_stats:
                mean, std = ver_stats[ver]
                c = ws.cell(row=dr, column=col, value=_fmt(mean, std))
                c.font = DATA_FONT
            else:
                c = ws.cell(row=dr, column=col, value="N/A")
                c.font = NA_FONT
            c.fill = fill; c.alignment = CENTER; c.border = tb()

        m, s = task_sr_stats
        if m is not None:
            c = ws.cell(row=dr, column=COL_MEAN, value=_fmt(m, s))
            c.font = BOLD_DATA_F
        else:
            c = ws.cell(row=dr, column=COL_MEAN, value="N/A")
            c.font = NA_FONT
        c.fill = fill; c.alignment = CENTER; c.border = tb()

        if var_cmds:
            collected_cmds = var_cmds

    # ── Variation commands row ───────────────────────────────────────────────
    cr = hr + 1 + len(models_data)
    ws.row_dimensions[cr].height = 28

    c = ws.cell(row=cr, column=COL_MODEL, value="Variation commands:")
    c.font = Font(name="Calibri", size=9, bold=True, italic=True, color="444444")
    c.fill = CMD_FILL; c.alignment = LEFT_WRAP; c.border = tb()

    for j, ver in enumerate(versions):
        cmd = collected_cmds.get(ver, "—")
        c = ws.cell(row=cr, column=COL_V0 + j, value=cmd)
        c.font = CMD_FONT; c.fill = CMD_FILL
        c.alignment = CENTER_WRAP; c.border = tb()

    c = ws.cell(row=cr, column=COL_MEAN, value="")
    c.fill = CMD_FILL; c.border = tb()

    return cr + 2  # 1-row gap before next table


def build_excel(output_path):
    wb = Workbook()
    wb.remove(wb.active)

    for level in LEVELS:
        ll = LEVEL_LABELS[level]
        ws = wb.create_sheet(title=ll)

        ws.column_dimensions["A"].width = 2   # left margin
        ws.column_dimensions["B"].width = 26  # model names

        current_row = 2

        for task_id in range(N_TASKS):
            models_data    = []
            task_name_used = None

            for model_name, model_dir in MODELS:
                t_name, v_stats, ts_stats, var_cmds = collect_task_data(
                    model_dir, level, task_id)
                if task_name_used is None and t_name:
                    task_name_used = t_name
                models_data.append((model_name, v_stats, ts_stats, var_cmds))

            if task_name_used is None:
                task_name_used = f"Task {task_id}"

            all_vers = set()
            for _, vs, _, _ in models_data:
                all_vers.update(vs.keys())

            if not all_vers:
                print(f"  [{ll} Task {task_id}] No data found — skipping.")
                continue

            versions  = [v for v in CANON_VERSIONS if v in all_vers]
            versions += sorted(v for v in all_vers if v not in CANON_VERSIONS)

            # Set version column widths (idempotent)
            for i in range(len(versions)):
                ws.column_dimensions[get_column_letter(3 + i)].width = 22
            ws.column_dimensions[get_column_letter(3 + len(versions))].width = 18

            print(f"  [{ll} Task {task_id}] \"{task_name_used}\" "
                  f"— versions: {versions}")

            current_row = write_task_table(
                ws, current_row, ll, task_id,
                task_name_used, models_data, versions)

        # Footer
        c = ws.cell(row=current_row + 1, column=2,
                    value=(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  "
                           "Source: InternVLA-M1 syntactic variation eval — libero_goal"))
        c.font = FOOTER_FONT

    wb.save(output_path)
    print(f"\nExcel saved → {output_path}")


# ──────────────────────────────────────────────────────────────────────────────
# MAIN
# ──────────────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(
        description="Generate per-task syntactic evaluation tables in Excel.")
    parser.add_argument("--output_xlsx", default="task_tables.xlsx",
                        help="Output Excel file path.")
    args = parser.parse_args()

    print("Generating task tables...\n")
    for model_name, model_dir in MODELS:
        print(f"  Model : {model_name}")
        print(f"  Dir   : {model_dir}")
    print()

    build_excel(args.output_xlsx)
    print("Done.")


if __name__ == "__main__":
    main()