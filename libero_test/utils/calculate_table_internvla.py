"""
calculate_table_internvla.py  — Sintattica DEFAULT/L1/L2/L3
"""

import re, math, argparse, os, glob
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


# ─── TASK ORDER ────────────────────────────────────────────────────────────────

def get_task_order():
    return [
        "Open the middle layer of the drawer",
        "Put the bowl on the stove",
        "Put the wine bottle on the top of the cabinet",
        "Open the top layer of the drawer and put the bowl inside",
        "Put the bowl on the top of the cabinet",
        "Push the plate to the front of the stove",
        "Put the cream cheese in the bowl",
        "Turn on the stove",
        "Put the bowl on the plate",
        "Put the wine bottle on the rack",
    ]


def get_variation_mapping():
    raw = {
        "open the middle layer of the drawer":                             "Open the middle layer of the drawer",
        "pull the middle layer of the drawer":                             "Open the middle layer of the drawer",
        "the middle layer of the drawer needs to be opened":               "Open the middle layer of the drawer",
        "open the layer of the drawer located between the top and bottom": "Open the middle layer of the drawer",

        "put the bowl on the stove":                                       "Put the bowl on the stove",
        "set the bowl on the stove":                                       "Put the bowl on the stove",
        "the stove needs to have the bowl on it":                          "Put the bowl on the stove",
        "put the object between the wine bottle and the cream cheese on the stove": "Put the bowl on the stove",

        "put the wine bottle on top of the cabinet":                       "Put the wine bottle on the top of the cabinet",
        "put the wine bottle on the top of the cabinet":                   "Put the wine bottle on the top of the cabinet",
        "place the wine bottle on the top of the cabinet":                 "Put the wine bottle on the top of the cabinet",
        "top of the cabinet needs to have the wine bottle on it":          "Put the wine bottle on the top of the cabinet",
        "put the object behind the bowl on the top of the cabinet":        "Put the wine bottle on the top of the cabinet",

        "open the top drawer and put the bowl inside":                     "Open the top layer of the drawer and put the bowl inside",
        "open the top layer of the drawer and put the bowl inside":        "Open the top layer of the drawer and put the bowl inside",
        "pull the top layer of the drawer and place the bowl inside":      "Open the top layer of the drawer and put the bowl inside",
        "pull the top layer of the drawer and put the bowl inside":        "Open the top layer of the drawer and put the bowl inside",
        "the top layer of the drawer needs to be opened and the bowl needs to be put inside": "Open the top layer of the drawer and put the bowl inside",
        "open the top layer of the drawer and put the object between the plate and the cream cheese inside": "Open the top layer of the drawer and put the bowl inside",

        "put the bowl on top of the cabinet":                              "Put the bowl on the top of the cabinet",
        "put the bowl on the top of the cabinet":                          "Put the bowl on the top of the cabinet",
        "place the bowl on the top of the cabinet":                        "Put the bowl on the top of the cabinet",
        "the top of the cabinet needs to have the bowl on it":             "Put the bowl on the top of the cabinet",
        "put the object between the wine bottle and the cream cheese on the top of the cabinet": "Put the bowl on the top of the cabinet",

        "push the plate to the front of the stove":                        "Push the plate to the front of the stove",
        "move the plate to the front of the stove":                        "Push the plate to the front of the stove",
        "the space in front of the stove needs to have the plate in it":   "Push the plate to the front of the stove",
        "push the object in front of the drawer to the front of the stove": "Push the plate to the front of the stove",

        "put the cream cheese in the bowl":                                "Put the cream cheese in the bowl",
        "put the cream cheese on the bowl":                                "Put the cream cheese in the bowl",
        "place the cream cheese on the bowl":                              "Put the cream cheese in the bowl",
        "place the cream cheese in the bowl":                              "Put the cream cheese in the bowl",
        "the cream cheese needs to be put on the bowl":                    "Put the cream cheese in the bowl",
        "put the object in front of the stove on the bowl":                "Put the cream cheese in the bowl",

        "turn on the stove":                                               "Turn on the stove",
        "switch on the stove":                                             "Turn on the stove",
        "the stove needs to be turned on":                                 "Turn on the stove",
        "turn on the object behind the cream cheese":                      "Turn on the stove",

        "put the bowl on the plate":                                       "Put the bowl on the plate",
        "place the bowl on the plate":                                     "Put the bowl on the plate",
        "the plate needs to have the bowl on it":                          "Put the bowl on the plate",
        "put the object between the wine bottle and the cream cheese on the plate": "Put the bowl on the plate",

        "put the wine bottle on the rack":                                 "Put the wine bottle on the rack",
        "place the wine bottle on the rack":                               "Put the wine bottle on the rack",
        "the rack needs to be filled with the wine bottle in it":          "Put the wine bottle on the rack",
        "put the object behind the bowl on the rack":                      "Put the wine bottle on the rack",
    }
    return {k.lower(): v for k, v in raw.items()}


# ─── PARSER ────────────────────────────────────────────────────────────────────

def parse_eval_log(filepath: str, num_trials: int = 50) -> list:
    """
    Parsa un file sintattica.

    Righe chiave:
      TASK X/10                    ← blocco task (senza "Task Composition")
      Original:  <testo>
      Variation [L3]: <testo>      ← assente per DEFAULT
      Task SR: 0.8800              ← DECIMALE (0.0–1.0) → moltiplichiamo ×100
    """
    tasks, current = [], {}

    with open(filepath, "r", errors="replace") as f:
        lines = [l.rstrip("\r\n") for l in f]

    for line in lines:
        # Blocco task sintattica: "TASK X/10" senza "(Task Composition"
        m = re.match(r"^TASK\s+(\d+)/\d+\s*$", line)
        if m:
            if current and "task_sr" in current:
                tasks.append(current)
            current = {"task_id": int(m.group(1))}
            continue

        if line.startswith("Original:") and current is not None:
            current["original"] = line.split("Original:", 1)[1].strip()
            continue

        m2 = re.match(r"^Variation \[(\w+)\]:\s*(.+)", line)
        if m2 and current is not None:
            current["variation_level"] = m2.group(1).upper()
            current["variation"]       = m2.group(2).strip()
            continue

        # "Task SR: 0.8800"  (decimale, nessuna parentesi)
        m3 = re.match(r"^Task SR:\s*([0-9.]+)\s*$", line)
        if m3 and current:
            sr_decimal          = float(m3.group(1))          # es. 0.8800
            current["task_sr"]  = sr_decimal * 100            # → 88.0  (percentuale)
            current["episodes"] = num_trials
            current["successes"] = int(round(sr_decimal * num_trials))  # 44

    if current and "task_sr" in current:
        tasks.append(current)

    tasks.sort(key=lambda t: t.get("task_id", 0))
    return tasks


def merge_eval_files(filepaths: list, num_trials: int = 50) -> tuple:
    """Fonde più file dello stesso seed (subset di task diversi)."""
    task_rates = defaultdict(list)   # key (original.lower()) → [pct, pct, ...]
    task_eps, task_orig, task_var = {}, {}, {}
    all_keys = []

    for fp in filepaths:
        for t in parse_eval_log(fp, num_trials):
            orig = t.get("original", "")
            key  = orig.lower()
            var  = t.get("variation", orig)
            if key not in task_rates:
                all_keys.append(key)
                task_orig[key] = orig
                task_var[key]  = var
            task_rates[key].append(t["task_sr"])   # già in %
            task_eps[key] = num_trials

    merged = {k: sum(v) / len(v) for k, v in task_rates.items()}
    return merged, task_eps, all_keys, task_orig, task_var


# ─── EXCEL ─────────────────────────────────────────────────────────────────────

def write_excel_syntactic(output_xlsx, txt_files_by_seed,
                           level="l3", num_trials=50, model_name="InternVLA-M1"):

    print("\n[INFO] Parsing log files (syntactic)...")
    all_merged = []
    for seed_idx in range(3):
        fps = txt_files_by_seed[seed_idx]
        print(f"  Seed {seed_idx}: {len(fps)} file(s)")
        merged = merge_eval_files(fps, num_trials)
        all_merged.append(merged)
        print(f"    → {len(merged[0])} task trovati")

    variation_to_original = get_variation_mapping()
    fixed_order           = get_task_order()

    # Mappa: original cased → log key
    orig_to_log_key = {}
    for seed_data in all_merged:
        merged_rates, _, all_keys, task_orig, task_var = seed_data
        for key in all_keys:
            mapped = variation_to_original.get(key)
            if mapped and mapped not in orig_to_log_key:
                orig_to_log_key[mapped] = key
    for orig in fixed_order:
        if orig not in orig_to_log_key:
            orig_to_log_key[orig] = orig.lower()

    wb = Workbook()
    ws = wb.active
    level_label = "DEFAULT" if level == "default" else f"L{level[-1].upper()}"
    ws.title = f"Syntactic {level_label}"

    headers = [
        "N°", "Original Task Command", "Variation Task Command",
        "SR% — Seed 0", "Comp. Seed 0",
        "SR% — Seed 1", "Comp. Seed 1",
        "SR% — Seed 2", "Comp. Seed 2",
        "Mean SR% ± Std%", "Mean Comp.",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font      = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    all_seed_rates = [[], [], []]
    all_seed_comps = [[], [], []]

    for task_num, orig_task in enumerate(fixed_order, start=1):
        log_key   = orig_to_log_key.get(orig_task, orig_task.lower())
        variation = None
        seed_rates, seed_comps = [], []

        for seed_idx, (merged_rates, task_eps, _, task_orig, task_var) in enumerate(all_merged):
            rate = merged_rates.get(log_key, float("nan"))   # già in %
            ep   = task_eps.get(log_key, num_trials)

            if not math.isnan(rate):
                succ = int(round(rate / 100 * ep))           # % → decimale → successi
            else:
                succ = 0

            if variation is None and log_key in task_var:
                v = task_var[log_key]
                if v.lower() != orig_task.lower():
                    variation = v

            seed_rates.append(rate)
            seed_comps.append(f"{succ}/{ep}")
            if not math.isnan(rate):
                all_seed_rates[seed_idx].append(rate)
                all_seed_comps[seed_idx].append((succ, ep))

        var_display = variation if variation else orig_task

        valid = [r for r in seed_rates if not math.isnan(r)]
        if valid:
            mean_r = sum(valid) / len(valid)
            std_r  = math.sqrt(sum((r - mean_r)**2 for r in valid) / max(len(valid)-1, 1))
            mean_display = f"{mean_r:.1f}% ± {std_r:.1f}%"
            avg_succ     = int(round(mean_r / 100 * num_trials))
        else:
            mean_display, avg_succ = "N/A", 0

        ws.append([
            task_num, orig_task, var_display,
            f"{seed_rates[0]:.1f}%" if not math.isnan(seed_rates[0]) else "N/A",
            seed_comps[0],
            f"{seed_rates[1]:.1f}%" if not math.isnan(seed_rates[1]) else "N/A",
            seed_comps[1],
            f"{seed_rates[2]:.1f}%" if not math.isnan(seed_rates[2]) else "N/A",
            seed_comps[2],
            mean_display,
            f"{avg_succ}/{num_trials}",
        ])

    # ── Riga finale ──
    final_row = ["", "Mean SR% ± Std%", ""]
    for seed_idx in range(3):
        rates = all_seed_rates[seed_idx]
        if rates:
            m = sum(rates) / len(rates)
            s = math.sqrt(sum((r-m)**2 for r in rates) / max(len(rates)-1, 1))
            ts = sum(c[0] for c in all_seed_comps[seed_idx])
            te = sum(c[1] for c in all_seed_comps[seed_idx])
            final_row.extend([f"{m:.2f}% ± {s:.2f}%", f"{ts}/{te}"])
        else:
            final_row.extend(["N/A", "0/0"])

    seed_means = [sum(all_seed_rates[i])/len(all_seed_rates[i])
                  for i in range(3) if all_seed_rates[i]]
    if seed_means:
        gm = sum(seed_means) / len(seed_means)
        gs = math.sqrt(sum((m-gm)**2 for m in seed_means) / max(len(seed_means)-1, 1))
        gt_s = sum(sum(c[0] for c in all_seed_comps[i]) for i in range(3))
        gt_e = sum(sum(c[1] for c in all_seed_comps[i]) for i in range(3))
        final_row.extend([f"{gm:.2f}% ± {gs:.2f}%", f"{gt_s}/{gt_e}"])
    else:
        final_row.extend(["N/A", "0/0"])

    ws.append(final_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    _autowidth(ws)
    ws.row_dimensions[1].height = 30
    wb.save(output_xlsx)
    print(f"\n[OK] Salvato: {output_xlsx}")


# ─── AUTO-DETECT ───────────────────────────────────────────────────────────────

def find_eval_files_by_seed(base_dir, level, pattern_prefix=None):
    if pattern_prefix is None:
        pattern_prefix = "EVAL-libero_goal-internvla_m1"
    result = {0: [], 1: [], 2: []}
    for seed_idx in range(3):
        matches = sorted(set(
            m for pat in [
                os.path.join(base_dir, f"{pattern_prefix}*seed{seed_idx}*-{level}.txt"),
                os.path.join(base_dir, f"{pattern_prefix}*seed_{seed_idx}*-{level}.txt"),
            ] for m in glob.glob(pat)
        ))
        result[seed_idx] = matches
        print(f"  {'✓' if matches else '✗'} Seed {seed_idx}: {len(matches)} file(s)")
    return result if all(result.values()) else None


def _autowidth(ws):
    for col in ws.columns:
        letter  = col[0].column_letter
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[letter].width = min(max_len + 2, 65)


# ─── MAIN ──────────────────────────────────────────────────────────────────────

def main():
    p = argparse.ArgumentParser()
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--txt_dir")
    g.add_argument("--manual", action="store_true")
    p.add_argument("--level",      default="l3", choices=["default","l1","l2","l3"])
    p.add_argument("--pattern",    default=None)
    p.add_argument("--model_name", default="InternVLA-M1")
    p.add_argument("--num_trials", type=int, default=50)
    p.add_argument("--seed0", nargs="+")
    p.add_argument("--seed1", nargs="+")
    p.add_argument("--seed2", nargs="+")
    p.add_argument("output_xlsx")
    args = p.parse_args()

    if args.manual:
        if not (args.seed0 and args.seed1 and args.seed2):
            exit("[ERROR] --manual richiede --seed0 --seed1 --seed2")
        files = {0: args.seed0, 1: args.seed1, 2: args.seed2}
    else:
        files = find_eval_files_by_seed(args.txt_dir, args.level, args.pattern)
        if files is None:
            exit("[ERROR] File mancanti per alcuni seed!")

    write_excel_syntactic(args.output_xlsx, files,
                          level=args.level,
                          num_trials=args.num_trials,
                          model_name=args.model_name)

if __name__ == "__main__":
    main()